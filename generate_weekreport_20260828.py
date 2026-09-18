#!/usr/bin/env python3
"""Build bilingual Week Report 2026-08-28 (.docx + .pdf) from the editable txt source.

Usage:
    python3 generate_weekreport_20260828.py

Reads:  uncoveredDocs/V20260828/weekReport_20260828_content.txt
Writes: uncoveredDocs/V20260828/weekReport_20260828_Bilingual_V20260828.docx
        uncoveredDocs/V20260828/weekReport_20260828_Bilingual_V20260828.pdf (via pandoc)

Txt format (one instruction per line):
    TITLE_CN / TITLE_EN   report title
    META_CN / META_EN     header line
    GREET_CN / GREET_EN   greeting
    H_CN / H_EN           section heading
    P_CN / P_EN           paragraph
    PHOTOS2: f1 | cap1_cn | cap1_en || f2 | cap2_cn | cap2_en   (side by side)
    PHOTO: file | caption_cn | caption_en | width_inches(optional)
    LINK: url
    SIGN_CN / SIGN_EN     closing / signature
    # comment
"""

import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIR = Path(__file__).parent / "uncoveredDocs" / "V20260828"
PHOTOS = DIR / "photos"
SRC = DIR / "weekReport_20260828_content.txt"
DOCX_OUT = DIR / "weekReport_20260828_Bilingual_V20260828.docx"
PDF_OUT = DIR / "weekReport_20260828_Bilingual_V20260828.pdf"


def img(name):
    """Resolve a photo: photos/ subfolder first, then the report folder."""
    p = PHOTOS / name
    return p if p.exists() else DIR / name


def add_para(doc, text, align=None, bold=False, size_pt=None, space_after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    if size_pt:
        run.font.size = Pt(size_pt)
    return p


def add_photo_cell(cell, img_path, caption_cn, caption_en, width_in):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Inches(width_in))
    for caption in (caption_cn, caption_en):
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.font.size = Pt(8)


def build_docx():
    doc = Document()
    for raw in SRC.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line or line.startswith("#"):
            continue
        if ": " not in line:
            continue
        tag, _, value = line.partition(": ")

        if tag in ("TITLE_CN", "TITLE_EN"):
            add_para(doc, value, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)
        elif tag == "LETTERHEAD":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(DIR / value), width=Inches(4.2))
        elif tag in ("META_CN", "META_EN", "GREET_CN", "GREET_EN"):
            add_para(doc, value, space_after=6)
        elif tag in ("CLOSE_CN", "CLOSE_EN"):
            add_para(doc, value, space_after=6)
        elif tag in ("H_CN", "H_EN"):
            add_para(doc, value, bold=True, size_pt=12, space_after=6)
        elif tag in ("P_CN", "P_EN"):
            add_para(doc, value, space_after=6)
        elif tag == "LINK":
            add_para(doc, value, space_after=12)
        elif tag in ("SIGN_CN", "SIGN_EN"):
            add_para(doc, value, space_after=4)
        elif tag == "PHOTO":
            parts = [s.strip() for s in value.split("|")]
            filename, cap_cn, cap_en = parts[0], parts[1], parts[2]
            width = float(parts[3]) if len(parts) > 3 else 4.0
            doc.add_picture(str(img(filename)), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_para(doc, cap_cn, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=9)
            add_para(doc, cap_en, align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=9, space_after=10)
        elif tag == "PHOTOS2":
            sides = value.split("||")
            if len(sides) != 2:
                print(f"WARN: bad PHOTOS2 line: {line}", file=sys.stderr)
                continue
            specs = [[s.strip() for s in side.split("|")] for side in sides]
            table = doc.add_table(rows=1, cols=2)
            table.autofit = True
            for i, spec in enumerate(specs):
                add_photo_cell(table.rows[0].cells[i], img(spec[0]), spec[1], spec[2], 2.9)
            doc.add_paragraph()
        else:
            print(f"WARN: unknown tag {tag!r}, skipped", file=sys.stderr)

    doc.save(DOCX_OUT)
    print(f"saved {DOCX_OUT}")


def build_pdf():
    subprocess.run(
        ["pandoc", str(DOCX_OUT), "-o", str(PDF_OUT), "--pdf-engine=xelatex",
         "-V", "CJKmainfont=PingFang SC", "-V", "mainfont=Times New Roman",
         "-V", "geometry:margin=2.5cm"],
        check=True,
    )
    print(f"saved {PDF_OUT}")


if __name__ == "__main__":
    build_docx()
    build_pdf()
