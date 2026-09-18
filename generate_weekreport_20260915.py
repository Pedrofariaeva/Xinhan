#!/usr/bin/env python3
"""Build bilingual Week Report 2026-09-15 (.docx + .pdf) from the editable txt source.

Usage:
    python3 generate_weekreport_20260915.py

Reads:  uncoveredDocs/V20260915/weekReport_20260915_content.txt
Writes: uncoveredDocs/V20260915/weekReport_20260915_Bilingual_V20260915.docx
        uncoveredDocs/V20260915/weekReport_20260915_Bilingual_V20260915.pdf

Same txt format as 2026-08-28, plus tables:
    TABLE: captionCN | captionEN
    ROW: cell | cell | ...        first ROW is the header
    ENDTABLE
    NOTE: small print under a table
Inside paragraphs and cells: **bold**; inside cells: <br> = line break;
a cell starting with ✓ is shaded green, with ✗ shaded red.
"""

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

DIR = Path(__file__).parent / "uncoveredDocs" / "V20260915"
PHOTOS = DIR / "photos"
SRC = DIR / "weekReport_20260915_content.txt"
DOCX_OUT = DIR / "weekReport_20260915_Bilingual_V20260915.docx"
PDF_OUT = DIR / "weekReport_20260915_Bilingual_V20260915.pdf"

HEADER_FILL = "1F3A5F"
GREEN_FILL = "E3F1E1"
RED_FILL = "F8E1DE"
STRIPE_FILL = "F4F6F8"
CJK_FONT = "PingFang SC"
LATIN_FONT = "Times New Roman"


def img(name):
    """Resolve a photo: photos/ subfolder first, then the report folder."""
    p = PHOTOS / name
    return p if p.exists() else DIR / name


def set_fonts(run):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), CJK_FONT)


def add_runs(p, text, bold=False, size_pt=None, color=None):
    """Add text to a paragraph, honouring **bold** and <br>."""
    for i, line in enumerate(text.split("<br>")):
        if i:
            p.add_run().add_break()
        for j, part in enumerate(re.split(r"\*\*", line)):
            if not part:
                continue
            run = p.add_run(part)
            run.bold = bold or (j % 2 == 1)
            if size_pt:
                run.font.size = Pt(size_pt)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            set_fonts(run)


def add_para(doc, text, align=None, bold=False, size_pt=None, space_after=None, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, text, bold=bold, size_pt=size_pt)
    if italic:
        for r in p.runs:
            r.italic = True
    return p


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, caption_cn, caption_en, rows):
    add_para(doc, caption_cn, bold=True, size_pt=10.5, space_after=0)
    add_para(doc, caption_en, bold=True, size_pt=10.5, space_after=4)
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            cell = table.rows[ri].cells[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if ri == 0:
                shade(cell, HEADER_FILL)
                add_runs(p, text, bold=True, size_pt=8.5, color="FFFFFF")
                continue
            if text.startswith("✓"):
                shade(cell, GREEN_FILL)
            elif text.startswith("✗"):
                shade(cell, RED_FILL)
            elif ri % 2 == 0:
                shade(cell, STRIPE_FILL)
            add_runs(p, text, bold=(ci == 0), size_pt=8.5)
    # repeat header row on each page
    trpr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK_FONT)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line or line.startswith("#"):
            continue
        if line == "ENDTABLE":
            continue
        if ": " not in line:
            continue
        tag, _, value = line.partition(": ")

        if tag in ("TITLE_CN", "TITLE_EN"):
            add_para(doc, value, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)
        elif tag == "LETTERHEAD":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(DIR / value), width=Inches(4.2))
        elif tag in ("META_CN", "META_EN", "GREET_CN", "GREET_EN", "CLOSE_CN", "CLOSE_EN"):
            add_para(doc, value, space_after=6)
        elif tag in ("H_CN", "H_EN"):
            add_para(doc, value, bold=True, size_pt=12, space_after=6)
        elif tag in ("P_CN", "P_EN"):
            add_para(doc, value, space_after=6)
        elif tag == "NOTE":
            add_para(doc, value, size_pt=8, space_after=12, italic=True)
        elif tag == "LINK":
            add_para(doc, value, space_after=12)
        elif tag in ("SIGN_CN", "SIGN_EN"):
            add_para(doc, value, space_after=4)
        elif tag == "TABLE":
            caps = [s.strip() for s in value.split("|")]
            rows = []
            while i < len(lines) and lines[i].strip() != "ENDTABLE":
                r = lines[i].strip()
                i += 1
                if r.startswith("ROW: "):
                    rows.append([c.strip() for c in r[5:].split("|")])
            i += 1
            add_table(doc, caps[0], caps[1] if len(caps) > 1 else "", rows)
        elif tag == "PHOTO":
            parts = [s.strip() for s in value.split("|")]
            width = float(parts[3]) if len(parts) > 3 else 4.0
            doc.add_picture(str(img(parts[0])), width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_para(doc, parts[1], align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=9)
            add_para(doc, parts[2], align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=9, space_after=10)
        else:
            print(f"WARN: unknown tag {tag!r}, skipped", file=sys.stderr)

    doc.save(DOCX_OUT)
    print(f"saved {DOCX_OUT}")


CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

HTML_CSS = """
@page { size: A4; margin: 18mm 18mm 18mm 18mm; }
body { font-family: 'Times New Roman', 'PingFang SC', serif; font-size: 10.5pt; line-height: 1.45; color: #111; }
.center { text-align: center; }
.title { font-size: 18pt; font-weight: bold; text-align: center; margin: 2pt 0; }
.meta { margin: 0 0 4pt 0; }
h2 { font-size: 12pt; margin: 14pt 0 5pt 0; page-break-after: avoid; }
p { margin: 0 0 7pt 0; text-align: justify; }
.cap { font-weight: bold; font-size: 10.5pt; margin: 12pt 0 4pt 0; page-break-after: avoid; }
table { width: 100%; border-collapse: collapse; font-size: 8.5pt; line-height: 1.3; page-break-inside: auto; }
tr { page-break-inside: avoid; }
th { background: #1F3A5F; color: #fff; text-align: left; padding: 4pt 5pt; border: 0.5pt solid #1F3A5F; }
td { padding: 3.5pt 5pt; border: 0.5pt solid #c9d1d9; vertical-align: top; }
td:first-child { font-weight: bold; }
tr:nth-child(even) td { background: #F4F6F8; }
td.ok { background: #E3F1E1 !important; }
td.no { background: #F8E1DE !important; }
.note { font-size: 8pt; font-style: italic; color: #444; margin: 4pt 0 12pt 0; }
"""


def html_text(text):
    import html as _h
    t = _h.escape(text).replace("&lt;br&gt;", "<br>")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)


def build_html():
    out = [f"<html><head><meta charset='utf-8'><style>{HTML_CSS}</style></head><body>"]
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line or line.startswith("#") or ": " not in line:
            continue
        tag, _, value = line.partition(": ")
        if tag == "LETTERHEAD":
            out.append(f"<div class='center'><img src='{(DIR / value).as_uri()}' style='width:4.2in'></div>")
        elif tag in ("TITLE_CN", "TITLE_EN"):
            out.append(f"<div class='title'>{html_text(value)}</div>")
        elif tag in ("META_CN", "META_EN", "GREET_CN", "GREET_EN", "CLOSE_CN", "CLOSE_EN", "SIGN_CN", "SIGN_EN"):
            out.append(f"<p class='meta'>{html_text(value)}</p>")
        elif tag in ("H_CN", "H_EN"):
            out.append(f"<h2>{html_text(value)}</h2>")
        elif tag in ("P_CN", "P_EN"):
            out.append(f"<p>{html_text(value)}</p>")
        elif tag == "NOTE":
            out.append(f"<p class='note'>{html_text(value)}</p>")
        elif tag == "LINK":
            out.append(f"<p>{html_text(value)}</p>")
        elif tag == "TABLE":
            caps = [s.strip() for s in value.split("|")]
            out.append("<div class='cap'>" + "<br>".join(html_text(c) for c in caps) + "</div><table>")
            first = True
            while i < len(lines) and lines[i].strip() != "ENDTABLE":
                r = lines[i].strip()
                i += 1
                if not r.startswith("ROW: "):
                    continue
                cells = [c.strip() for c in r[5:].split("|")]
                if first:
                    out.append("<thead><tr>" + "".join(f"<th>{html_text(c)}</th>" for c in cells) + "</tr></thead><tbody>")
                    first = False
                    continue
                tds = []
                for c in cells:
                    cls = " class='ok'" if c.startswith("✓") else " class='no'" if c.startswith("✗") else ""
                    tds.append(f"<td{cls}>{html_text(c)}</td>")
                out.append("<tr>" + "".join(tds) + "</tr>")
            i += 1
            out.append("</tbody></table>")
        elif tag == "PHOTO":
            parts = [s.strip() for s in value.split("|")]
            width = float(parts[3]) if len(parts) > 3 else 4.0
            out.append(f"<div class='center'><img src='{img(parts[0]).as_uri()}' style='width:{width}in'>"
                       f"<p class='center'>{html_text(parts[1])}<br>{html_text(parts[2])}</p></div>")
    out.append("</body></html>")
    return "\n".join(out)


def build_pdf():
    """Chrome keeps the table shading and fonts; pandoc is the fallback."""
    if Path(CHROME).exists():
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            page = Path(tmp) / "report.html"
            page.write_text(build_html(), encoding="utf-8")
            r = subprocess.run(
                [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                 "--allow-file-access-from-files", f"--print-to-pdf={PDF_OUT}", page.as_uri()],
                capture_output=True, text=True, timeout=180,
            )
        if r.returncode == 0 and PDF_OUT.exists() and PDF_OUT.stat().st_mtime >= DOCX_OUT.stat().st_mtime:
            print(f"saved {PDF_OUT}")
            return
        print("chrome failed, falling back to pandoc", file=sys.stderr)
    subprocess.run(
        ["pandoc", str(DOCX_OUT), "-o", str(PDF_OUT), "--pdf-engine=xelatex",
         "-V", f"CJKmainfont={CJK_FONT}", "-V", f"mainfont={LATIN_FONT}",
         "-V", "geometry:margin=2cm"],
        check=True,
    )
    print(f"saved {PDF_OUT}")


if __name__ == "__main__":
    build_docx()
    build_pdf()
