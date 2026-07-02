#!/usr/bin/env python3
"""Generate WPR2 bilingual document from MD files."""

import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_format(run, bold=False, size_pt=None, italic=False):
    run.bold = bold
    run.italic = italic
    if size_pt:
        run.font.size = Pt(size_pt)


def add_para(doc, text, align=None, bold=False, size_pt=None, space_before=None, space_after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_format(run, bold=bold, size_pt=size_pt)
    return p


def parse_md_sections(md_text):
    """Parse markdown into sections with heading and blocks."""
    lines = md_text.split('\n')
    sections = []
    current_heading = None
    current_blocks = []
    i = 0

    def flush():
        if current_heading is not None:
            sections.append({'heading': current_heading, 'blocks': current_blocks})

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Heading
        m = re.match(r'^(#{1,3})\s+(.*)', stripped)
        if m:
            flush()
            current_heading = m.group(2).strip()
            current_blocks = []
            i += 1
            continue

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == '---' or stripped == '***' or stripped == '___':
            i += 1
            continue

        # Table
        if '|' in stripped and not stripped.startswith('>') and not re.match(r'^(\d+\.|[*-])\s', stripped):
            # Check if next line is separator
            if i + 1 < len(lines) and re.match(r'^\|?[\s\-:|]+\|?$', lines[i + 1].strip()):
                table_lines = [stripped]
                i += 1
                while i < len(lines) and '|' in lines[i].strip():
                    table_lines.append(lines[i].strip())
                    i += 1
                current_blocks.append({'type': 'table', 'lines': table_lines})
                continue

        # Numbered or bullet list item
        m_list = re.match(r'^(\d+\.|[*-])\s+(.*)', stripped)
        if m_list:
            items = [m_list.group(2)]
            i += 1
            saw_blank = False
            while i < len(lines):
                s = lines[i].strip()
                if not s:
                    saw_blank = True
                    i += 1
                    continue
                if re.match(r'^(\d+\.|[*-])\s', s):
                    saw_blank = False
                    items.append(re.match(r'^(\d+\.|[*-])\s+(.*)', s).group(2))
                    i += 1
                elif s.startswith('#') or s.startswith('|') or s.startswith('>'):
                    break
                elif saw_blank:
                    # After a blank line, this is likely a new paragraph, not a list continuation
                    break
                else:
                    # continuation of current item (no blank line before)
                    items[-1] += ' ' + s
                    i += 1
            current_blocks.append({'type': 'list', 'items': items})
            continue

        # Blockquote
        if stripped.startswith('>'):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote_lines.append(lines[i].strip().lstrip('>').strip())
                i += 1
            current_blocks.append({'type': 'quote', 'text': '\n'.join(quote_lines)})
            continue

        # Regular paragraph
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                break
            if s.startswith('#') or s.startswith('|') or s.startswith('>') or re.match(r'^(\d+\.|[*-])\s', s) or s == '---':
                break
            para_lines.append(s)
            i += 1
        current_blocks.append({'type': 'paragraph', 'text': ' '.join(para_lines)})

    flush()
    return sections


def add_rich_text(p, text):
    """Add text with **bold** support."""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)


def add_block_to_doc(doc, block, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    if block['type'] == 'paragraph':
        p = doc.add_paragraph()
        p.alignment = align
        add_rich_text(p, block['text'])
    elif block['type'] == 'list':
        for item in block['items']:
            p = doc.add_paragraph(style='List Bullet')
            p.alignment = align
            add_rich_text(p, item)
    elif block['type'] == 'quote':
        p = doc.add_paragraph()
        p.alignment = align
        run = p.add_run(block['text'])
        run.italic = True
    elif block['type'] == 'table':
        lines = block['lines']
        if len(lines) < 2:
            return
        rows_data = []
        for line in lines:
            if re.match(r'^\|?[\s\-:|]+\|?$', line):
                continue
            cells = [c.strip() for c in line.split('|')]
            while cells and not cells[0]:
                cells = cells[1:]
            while cells and not cells[-1]:
                cells = cells[:-1]
            if cells:
                rows_data.append(cells)
        if not rows_data:
            return
        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.style = 'Table Grid'
        for ri, row in enumerate(rows_data):
            for ci, cell_text in enumerate(row):
                if ci < num_cols:
                    table.rows[ri].cells[ci].text = cell_text


def generate_wpr2():
    with open('uncoveredDocs/WPR20250421_2_Revised_EN.md', 'r', encoding='utf-8') as f:
        en_text = f.read()
    with open('uncoveredDocs/WPR20250421_2_CN.md', 'r', encoding='utf-8') as f:
        cn_text = f.read()

    en_sections = parse_md_sections(en_text)
    cn_sections = parse_md_sections(cn_text)

    # Skip first section if it's title/date/subject
    if en_sections and 'WPR' in en_sections[0]['heading']:
        en_sections = en_sections[1:]
    if cn_sections and 'WPR' in cn_sections[0]['heading']:
        cn_sections = cn_sections[1:]

    doc = Document()

    # Title
    add_para(doc, '工作进展报告 第2号', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)
    add_para(doc, 'Working Progress Report （WPR 2）', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)

    # Date and Subject
    add_para(doc, '日期： 2025年4月21日')
    add_para(doc, 'Date: 21 April 2025')
    add_para(doc, '主题： 崖州兽医院及兽医学院项目设计审查')
    add_para(doc, 'Subject: Design Review — Yazhou Veterinary Hospital & School Project')
    doc.add_paragraph()

    for cn_sec, en_sec in zip(cn_sections, en_sections):
        cn_heading = cn_sec['heading']
        en_heading = en_sec['heading']

        if en_heading == cn_heading:
            heading_text = cn_heading
        else:
            heading_text = f"{cn_heading} / {en_heading}"

        add_para(doc, heading_text, bold=True, size_pt=12, space_before=12, space_after=6)

        for block in cn_sec['blocks']:
            add_block_to_doc(doc, block)

        for block in en_sec['blocks']:
            add_block_to_doc(doc, block)

        doc.add_paragraph()

    output_path = 'uncoveredDocs/WPR20250421_2_Bilingual_V20260609.docx'
    doc.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    generate_wpr2()
