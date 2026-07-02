from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Open existing presentation
prs = Presentation('/Users/pedro/Documents/GitHub/Xinhan/lessons/friendship_presentation.pptx')

# Images path
img_dir = '/Users/pedro/Documents/GitHub/Xinhan/lessons/images/'

# Define image placements: (slide_index, image_file, left, top, width, height)
# Placed in corners to avoid overlapping text boxes
placements = [
    # Slide 1 (Title) - friends photo bottom right
    (0, 'friends_playing.jpg', 9.5, 4.2, 3.5, 2.2),
    
    # Slide 2 (Warm-up) - students top right
    (1, 'students_classroom.jpg', 10.0, 0.8, 2.8, 2.2),
    
    # Slide 3 (Types of Friends) - school activity bottom right
    (2, 'school_activity.jpg', 10.0, 5.0, 2.8, 2.0),
    
    # Slide 9 (Speaking Model) - chatting friends bottom right
    (8, 'chatting.jpg', 10.0, 5.0, 2.8, 2.2),
    
    # Slide 10 (Speaking Framework) - books top right
    (9, 'books_reading.jpg', 10.0, 0.8, 2.8, 1.9),
    
    # Slide 16 (Dialogue Making Friends) - friends playing bottom right
    (15, 'friends_playing.jpg', 10.0, 5.2, 2.8, 2.0),
    
    # Slide 18 (Story) - birthday bottom right
    (17, 'birthday.jpg', 10.0, 5.0, 2.8, 2.2),
    
    # Slide 20 (Homework) - students bottom right
    (19, 'students_classroom.jpg', 10.0, 5.2, 2.8, 2.2),
]

for slide_idx, img_file, left, top, width, height in placements:
    slide = prs.slides[slide_idx]
    img_path = img_dir + img_file
    try:
        pic = slide.shapes.add_picture(img_path, Inches(left), Inches(top), Inches(width), Inches(height))
        # Send to back so text stays on top
        spTree = slide.shapes._spTree
        sp = pic._element
        spTree.remove(sp)
        spTree.insert(2, sp)
    except Exception as e:
        print(f"Error adding {img_file} to slide {slide_idx+1}: {e}")

# Save
prs.save('/Users/pedro/Documents/GitHub/Xinhan/lessons/friendship_presentation.pptx')
print("Images added successfully!")
