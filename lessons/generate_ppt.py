from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

img_dir = '/Users/pedro/Documents/GitHub/Xinhan/lessons/images/'

def add_background(slide, r, g, b):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(r, g, b)

def add_textbox(slide, left, top, width, height, text, font_size, bold=False, italic=False, color=(0,0,0), align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    font = p.font
    font.size = Pt(font_size)
    font.bold = bold
    font.italic = italic
    font.name = 'Microsoft YaHei'
    font.color.rgb = RGBColor(*color)
    return txBox

def add_bullet_box(slide, left, top, width, height, items, font_size, color=(0,0,0)):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + item
        p.level = 0
        font = p.font
        font.size = Pt(font_size)
        font.name = 'Microsoft YaHei'
        font.color.rgb = RGBColor(*color)
    return txBox

def add_framed_image(slide, img_file, left, top, width, height):
    img_path = img_dir + img_file
    # White rounded frame (behind image)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left - 0.08), Inches(top - 0.08), Inches(width + 0.16), Inches(height + 0.16))
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.fill.background()
    frame.adjustments[0] = 0.08
    # Image (on top of frame so it's visible)
    pic = slide.shapes.add_picture(img_path, Inches(left), Inches(top), Inches(width), Inches(height))
    spTree = slide.shapes._spTree
    spTree.remove(frame._element)
    spTree.remove(pic._element)
    spTree.insert(3, frame._element)
    spTree.insert(4, pic._element)

# Color palettes (soft gradients simulated with solid colors)
COLORS = [
    (255, 154, 158),   # pink
    (161, 140, 209),   # purple
    (132, 250, 176),   # green
    (252, 203, 144),   # orange
    (224, 195, 252),   # lavender
    (255, 236, 210),   # cream
    (161, 196, 253),   # blue
    (212, 252, 121),   # lime
    (251, 194, 235),   # light pink
    (253, 203, 241),   # rose
    (137, 247, 254),   # cyan
    (253, 219, 146),   # yellow
    (168, 237, 234),   # mint
    (255, 154, 158),   # pink
    (102, 126, 234),   # dark blue
    (240, 147, 251),   # magenta
    (79, 172, 254),    # bright blue
    (67, 233, 123),    # green
    (250, 112, 154),   # coral
    (48, 207, 208),    # teal
]

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[0])
add_textbox(slide, 1, 2.5, 11.333, 1.5, "💖 Friendship", 60, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 4, 11.333, 0.8, "The Best Gift in Life", 32, color=(99, 110, 114), align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 5.2, 11.333, 0.6, "A 2-Hour Journey to Better Speaking & Writing", 22, color=(99, 110, 114), align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 5.8, 11.333, 0.5, "For Grade 7 • Spoken English • Test Prep", 18, color=(178, 190, 195), align=PP_ALIGN.CENTER)
# Bottom center image on title slide
add_framed_image(slide, 'friends_hanging.jpg', 4.4, 4.6, 4.5, 2.4)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "1 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 2: Warm-up
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[1])
add_textbox(slide, 1, 0.6, 11.333, 1, "🤔 Warm-Up Discussion", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2), Inches(1.8), Inches(9.333), Inches(4.5))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.2
box.line.fill.background()
add_bullet_box(slide, 2.5, 2.2, 8.333, 4, [
    "How many good friends do you have?",
    "What do you and your friends like to do together?",
    "Can you describe your best friend in 3 words?",
    "What makes a true friend?"
], 28, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "2 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 3: Types of Friends (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[2])
add_textbox(slide, 1, 0.6, 11.333, 1, "👥 Types of Friends", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
items = [
    "⭐  best friend",
    "🤝  close friend",
    "📚  classmate",
    "💻  online friend",
    "🏠  neighbor",
    "🏃  teammate",
    "✉️  pen pal",
    "👶  childhood friend"
]
# 2 cols on left, narrower to make room for image
for i, item in enumerate(items):
    row = i // 2
    col = i % 2
    x = 1.0 + col * 4.0
    y = 1.9 + row * 0.9
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(3.7), Inches(0.7))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.fill.fore_color.brightness = 0.1
    box.line.fill.background()
    add_textbox(slide, x + 0.2, y + 0.1, 3.3, 0.5, item, 22, color=(45, 52, 54))
# Image on right
add_framed_image(slide, 'school_kids.jpg', 9.0, 1.9, 3.5, 4.5)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "3 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 4: Personality 1 (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[3])
add_textbox(slide, 1, 0.6, 11.333, 1, "✨ Personality Adjectives (1)", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
items = [
    "😊  kind  善良的",
    "🆘  helpful  乐于助人的",
    "🤣  funny  有趣的",
    "🗣️  outgoing  外向的",
    "😇  honest  诚实的",
    "🛡️  loyal  忠诚的",
    "📖  hard-working  勤奋的",
    "💡  clever  聪明的"
]
for i, item in enumerate(items):
    row = i // 2
    col = i % 2
    x = 0.8 + col * 4.0
    y = 1.85 + row * 0.85
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(3.8), Inches(0.7))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.fill.fore_color.brightness = 0.1
    box.line.fill.background()
    add_textbox(slide, x + 0.2, y + 0.1, 3.4, 0.5, item, 20, color=(45, 52, 54))
add_framed_image(slide, 'girls_chatting.jpg', 9.0, 1.9, 3.5, 4.5)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "4 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 5: Personality 2 (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[4])
add_textbox(slide, 1, 0.6, 11.333, 1, "🌈 Personality Adjectives (2)", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
items = [
    "😌  patient  有耐心的",
    "❤️  caring  关心他人的",
    "🦁  brave  勇敢的",
    "🎨  creative  有创造力的",
    "🤫  shy / quiet  害羞的/安静的",
    "🎁  generous  大方的",
    "🧘  understanding  善解人意的",
    "⚽  sporty  爱运动的"
]
for i, item in enumerate(items):
    row = i // 2
    col = i % 2
    x = 0.8 + col * 4.0
    y = 1.85 + row * 0.85
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(3.8), Inches(0.7))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box.fill.fore_color.brightness = 0.1
    box.line.fill.background()
    add_textbox(slide, x + 0.2, y + 0.1, 3.4, 0.5, item, 20, color=(45, 52, 54))
add_framed_image(slide, 'girls_chatting.jpg', 9.0, 1.9, 3.5, 4.5)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "5 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 6: Key Phrases
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[5])
add_textbox(slide, 1, 0.6, 11.333, 1, "💬 Key Phrases About Friendship", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box1 = slide.shapes.add_shape(1, Inches(1.5), Inches(1.8), Inches(5), Inches(3.8))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(255, 255, 255)
box1.fill.fore_color.brightness = 0.1
box1.line.fill.background()
add_bullet_box(slide, 1.8, 2.1, 4.4, 3.5, [
    "get along well with  与…相处融洽",
    "have something in common  有共同之处",
    "be good at  擅长",
    "care about  关心"
], 24, color=(45, 52, 54))
box2 = slide.shapes.add_shape(1, Inches(6.833), Inches(1.8), Inches(5), Inches(3.8))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.1
box2.line.fill.background()
add_bullet_box(slide, 7.133, 2.1, 4.4, 3.5, [
    "help ... with ...  帮助…做…",
    "share ... with ...  与…分享",
    "trust each other  互相信任",
    "keep secrets  保守秘密"
], 24, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "6 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 7: Grammar Describing People
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[6])
add_textbox(slide, 1, 0.6, 11.333, 1, "📐 Grammar: Describing People", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2), Inches(1.7), Inches(9.333), Inches(4.8))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_textbox(slide, 2.5, 1.9, 8.333, 0.6, "Present Simple:  He / She + is / has + ...", 26, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 2.8, 2.6, 8, 3.8, [
    "She is tall and thin.",
    "He has short black hair.",
    "She wears glasses.",
    "He is very funny and outgoing.",
    "She is good at English."
], 26, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "7 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 8: Grammar Habits
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[7])
add_textbox(slide, 1, 0.5, 11.333, 1, "🔄 Present Simple vs. Present Continuous", 40, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box1 = slide.shapes.add_shape(1, Inches(1.5), Inches(1.6), Inches(5), Inches(3.5))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(255, 255, 255)
box1.fill.fore_color.brightness = 0.1
box1.line.fill.background()
add_textbox(slide, 1.8, 1.8, 4.4, 0.5, "Simple (habits)", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 2, 2.4, 4.2, 2.5, [
    "She always helps me.",
    "We usually chat after school.",
    "He often shares his snacks."
], 22, color=(45, 52, 54))
box2 = slide.shapes.add_shape(1, Inches(6.833), Inches(1.6), Inches(5), Inches(3.5))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.1
box2.line.fill.background()
add_textbox(slide, 7.133, 1.8, 4.4, 0.5, "Continuous (now)", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 7.333, 2.4, 4.2, 2.5, [
    "She is helping me now.",
    "We are chatting on WeChat.",
    "He is sharing his umbrella."
], 22, color=(45, 52, 54))
add_textbox(slide, 1, 5.5, 11.333, 0.6, "Use always / usually / often with Present Simple!", 22, color=(99, 110, 114), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "8 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 9: Speaking Model (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[8])
add_textbox(slide, 1, 0.6, 11.333, 1, "🎤 Speaking Model", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 1.4, 7.5, 0.5, "My Best Friend", 28, bold=True, color=(45, 52, 54))
box = slide.shapes.add_shape(1, Inches(0.8), Inches(1.9), Inches(8), Inches(4.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.1
box.line.color.rgb = RGBColor(232, 67, 147)
box.line.width = Pt(4)
add_textbox(slide, 1.1, 2.1, 7.4, 3.8, "Lin Mei is my best friend. We have been classmates since Grade 5. She is kind and hard-working. She has long hair and always wears a smile. She is good at English and often helps me with my homework. We both love playing badminton after school. Last week, she shared her umbrella with me when it rained. I think a good friend is someone who truly cares about you. Lin Mei is such a friend.", 20, color=(45, 52, 54))
add_framed_image(slide, 'girls_chatting.jpg', 9.2, 1.9, 3.5, 4.7)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "9 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 10: Speaking Framework (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[9])
add_textbox(slide, 1, 0.6, 11.333, 1, "📝 Speaking Framework", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(8.2), Inches(4.5))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_textbox(slide, 1.1, 1.8, 7.6, 0.5, "How to Describe Your Best Friend (1 minute)", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 1.3, 2.4, 7.2, 3.8, [
    "1. Basic info: name, age, how you met",
    "2. Appearance: height, hair, glasses, clothes",
    "3. Personality: kind, funny, helpful...",
    "4. Hobbies: good at, likes to...",
    "5. A story: what you did together",
    "6. Why you like him/her"
], 23, color=(45, 52, 54))
add_framed_image(slide, 'books_friends.jpg', 9.2, 1.9, 3.5, 4.2)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "10 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 11: Speaking Practice
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[10])
add_textbox(slide, 1, 0.6, 11.333, 1, "🗣️ Speaking Practice: Pair Work", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box1 = slide.shapes.add_shape(1, Inches(1.5), Inches(1.7), Inches(5), Inches(4))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(255, 255, 255)
box1.fill.fore_color.brightness = 0.1
box1.line.fill.background()
add_textbox(slide, 1.8, 1.9, 4.4, 0.5, "Student A", 26, bold=True, color=(9, 132, 227))
add_textbox(slide, 2, 2.5, 4.2, 2.5, "Describe your best friend. Use at least 5 new words from today's lesson.", 20, color=(45, 52, 54))
box2 = slide.shapes.add_shape(1, Inches(6.833), Inches(1.7), Inches(5), Inches(4))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.1
box2.line.fill.background()
add_textbox(slide, 7.133, 1.9, 4.4, 0.5, "Student B", 26, bold=True, color=(232, 67, 147))
add_textbox(slide, 7.333, 2.5, 4.2, 3, "Listen and ask one question:\n\n• How long have you known each other?\n• What do you usually do together?\n• What makes him/her special?", 19, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "11 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 12: Writing Structure
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[11])
add_textbox(slide, 1, 0.6, 11.333, 1, "✍️ Writing Structure", 48, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2), Inches(1.6), Inches(9.333), Inches(4.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_textbox(slide, 2.3, 1.8, 8.333, 0.5, "A Good Paragraph About a Friend", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 2.5, 2.4, 8, 3.5, [
    "Topic sentence: Who is this person?",
    "Detail 1: Appearance (looks)",
    "Detail 2: Personality & hobbies",
    "Detail 3: A story or memory",
    "Closing sentence: Why is this friendship important?"
], 24, color=(45, 52, 54))
add_textbox(slide, 1, 6, 11.333, 0.5, "Target length: 60–80 words (typical test requirement)", 20, color=(99, 110, 114), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "12 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 13: Guided Writing
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[12])
add_textbox(slide, 1, 0.6, 11.333, 1, "✏️ Guided Writing Practice", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(1.5), Inches(1.7), Inches(10.333), Inches(3.8))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.1
box.line.color.rgb = RGBColor(232, 67, 147)
box.line.width = Pt(3)
add_textbox(slide, 1.8, 2, 9.7, 3.5, "My best friend is __________. We met __________. He/She is __________ and __________. He/She has __________. We both like __________. Last week, __________. I feel __________ because __________.", 22, color=(45, 52, 54))
add_textbox(slide, 1, 5.8, 11.333, 0.5, "Now rewrite it as a full paragraph!", 24, bold=True, color=(232, 67, 147), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "13 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 14: Test Prep My Best Friend
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[13])
add_textbox(slide, 1, 0.6, 11.333, 1, "🎯 Test Prep: \"My Best Friend\"", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2), Inches(1.7), Inches(9.333), Inches(4))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_textbox(slide, 2.3, 1.9, 8.333, 0.5, "Common Writing Prompt (60–80 words)", 24, bold=True, color=(45, 52, 54))
add_textbox(slide, 2.5, 2.5, 8, 0.5, "Write about your best friend. Include:", 22, color=(45, 52, 54))
add_bullet_box(slide, 2.7, 3, 7.5, 2.5, [
    "Introduction: name & relationship",
    "Appearance & personality",
    "Hobbies & abilities",
    "Why you are good friends"
], 22, color=(45, 52, 54))
box2 = slide.shapes.add_shape(1, Inches(2.5), Inches(5.1), Inches(8.333), Inches(0.8))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.05
box2.line.fill.background()
add_textbox(slide, 2.7, 5.25, 8, 0.5, "Tip: Use connecting words: and, but, because, so, when, if", 20, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "14 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 15: Test Prep Narrative
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[14])
add_textbox(slide, 1, 0.6, 11.333, 1, "📝 Test Prep: Narrative Writing", 44, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2), Inches(1.7), Inches(9.333), Inches(4.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_textbox(slide, 2.3, 1.9, 8.333, 0.5, "\"An Unforgettable Day with My Friend\"", 24, bold=True, color=(45, 52, 54))
add_textbox(slide, 2.5, 2.5, 8, 0.5, "Key elements to include:", 22, color=(45, 52, 54))
add_bullet_box(slide, 2.7, 3, 7.5, 2.5, [
    "When & Where – set the scene",
    "What happened – the event",
    "How you helped each other",
    "What you learned about friendship"
], 22, color=(45, 52, 54))
add_textbox(slide, 1, 6.2, 11.333, 0.5, "Use past tense: was, were, did, went, helped, shared...", 22, color=(251, 194, 235), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "15 / 20", 14, color=(255, 255, 255), align=PP_ALIGN.RIGHT)

# Slide 16: Dialogue Making Friends (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[15])
add_textbox(slide, 1, 0.6, 11.333, 1, "🎭 Dialogue: Making Friends", 44, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(0.6), Inches(1.7), Inches(8.4), Inches(4.4))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.05
box.line.fill.background()
lines = [
    ("A:", "Hi! I'm Lily. What's your name?", (9, 132, 227)),
    ("B:", "I'm Mei. Nice to meet you!", (232, 67, 147)),
    ("A:", "Nice to meet you too! Do you like playing basketball?", (9, 132, 227)),
    ("B:", "Yes, I love it! I play every weekend.", (232, 67, 147)),
    ("A:", "Great! We can play together after school.", (9, 132, 227)),
    ("B:", "That's wonderful! Let's be friends.", (232, 67, 147))
]
for i, (sp, text, color) in enumerate(lines):
    y = 1.9 + i * 0.55
    add_textbox(slide, 0.9, y, 0.5, 0.4, sp, 21, bold=True, color=color)
    add_textbox(slide, 1.5, y, 7.2, 0.4, text, 21, color=(45, 52, 54))
add_framed_image(slide, 'friends_hanging.jpg', 9.2, 2.0, 3.5, 4.0)
add_textbox(slide, 1, 6.2, 8, 0.5, "Practice with a partner! Try changing the hobby.", 20, color=(251, 194, 235), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "16 / 20", 14, color=(255, 255, 255), align=PP_ALIGN.RIGHT)

# Slide 17: Dialogue Helping
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[16])
add_textbox(slide, 1, 0.6, 11.333, 1, "🎭 Dialogue: Helping a Friend", 44, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(2.5), Inches(1.7), Inches(8.333), Inches(4.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.05
box.line.fill.background()
lines = [
    ("A:", "You look sad. What's wrong?", (9, 132, 227)),
    ("B:", "I failed my English test. I feel terrible.", (232, 67, 147)),
    ("A:", "Don't worry. I can help you with vocabulary.", (9, 132, 227)),
    ("B:", "Really? That's so kind of you!", (232, 67, 147)),
    ("A:", "We can study together every Saturday.", (9, 132, 227)),
    ("B:", "Thank you! You are a true friend.", (232, 67, 147))
]
for i, (sp, text, color) in enumerate(lines):
    y = 1.9 + i * 0.55
    add_textbox(slide, 2.8, y, 0.5, 0.4, sp, 21, bold=True, color=color)
    add_textbox(slide, 3.5, y, 7, 0.4, text, 21, color=(45, 52, 54))
add_textbox(slide, 1, 6.2, 11.333, 0.5, "Key phrases: Don't worry. / I can help you with... / That's so kind of you!", 18, color=(194, 233, 251), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "17 / 20", 14, color=(255, 255, 255), align=PP_ALIGN.RIGHT)

# Slide 18: Story (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[17])
add_textbox(slide, 1, 0.6, 11.333, 1, "📖 A Story About Friendship", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(0.6), Inches(1.7), Inches(8.4), Inches(4.5))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.05
box.line.color.rgb = RGBColor(232, 67, 147)
box.line.width = Pt(3)
add_textbox(slide, 0.9, 1.9, 8, 0.5, "Sharing an Umbrella ☂️", 26, bold=True, color=(45, 52, 54))
add_textbox(slide, 0.9, 2.5, 7.9, 3.5, "Last Friday, it rained heavily after school. I forgot my umbrella. My friend Li Hua saw me standing at the school gate. She walked over and said, \"Let's share my umbrella!\" Her umbrella was small, but we walked home together. We both got a little wet, but we laughed all the way. That day, I learned: a true friend is someone who gets wet with you.", 19, color=(45, 52, 54))
add_framed_image(slide, 'birthday_party.jpg', 9.2, 1.9, 3.5, 4.3)
box2 = slide.shapes.add_shape(1, Inches(1.0), Inches(5.4), Inches(7.8), Inches(0.8))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.1
box2.line.fill.background()
add_textbox(slide, 1.2, 5.55, 7.4, 0.5, "Discuss: What did Li Hua do? Why was it kind? Have you ever helped a friend?", 18, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "18 / 20", 14, color=(99, 110, 114), align=PP_ALIGN.RIGHT)

# Slide 19: Review
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[18])
add_textbox(slide, 1, 0.6, 11.333, 1, "🔑 Review & Key Takeaways", 44, bold=True, color=(45, 52, 54), align=PP_ALIGN.CENTER)
box1 = slide.shapes.add_shape(1, Inches(1.5), Inches(1.7), Inches(5), Inches(4.2))
box1.fill.solid()
box1.fill.fore_color.rgb = RGBColor(255, 255, 255)
box1.fill.fore_color.brightness = 0.1
box1.line.fill.background()
add_textbox(slide, 1.8, 1.9, 4.4, 0.5, "Key Words", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 2, 2.5, 4.2, 3, [
    "loyal, honest, kind",
    "helpful, funny, caring",
    "get along well with",
    "have something in common",
    "care about, trust",
    "share ... with ..."
], 20, color=(45, 52, 54))
box2 = slide.shapes.add_shape(1, Inches(6.833), Inches(1.7), Inches(5), Inches(4.2))
box2.fill.solid()
box2.fill.fore_color.rgb = RGBColor(255, 255, 255)
box2.fill.fore_color.brightness = 0.1
box2.line.fill.background()
add_textbox(slide, 7.133, 1.9, 4.4, 0.5, "Key Structures", 24, bold=True, color=(45, 52, 54))
add_bullet_box(slide, 7.333, 2.5, 4.2, 3, [
    "He/She is + adj.",
    "He/She has + noun",
    "We like to + verb",
    "We are good at + v-ing",
    "A true friend is someone who..."
], 20, color=(45, 52, 54))
add_textbox(slide, 11, 6.8, 1.5, 0.3, "19 / 20", 14, color=(99, 52, 54), align=PP_ALIGN.RIGHT)

# Slide 20: Homework (with image on right)
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide, *COLORS[19])
add_textbox(slide, 1, 0.6, 11.333, 1, "📝 Homework & Next Steps", 44, bold=True, color=(255, 255, 255), align=PP_ALIGN.CENTER)
box = slide.shapes.add_shape(1, Inches(0.6), Inches(1.7), Inches(8.4), Inches(4.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(255, 255, 255)
box.fill.fore_color.brightness = 0.15
box.line.fill.background()
add_bullet_box(slide, 0.9, 1.9, 7.8, 3.8, [
    "Write a paragraph (80 words) about your best friend.",
    "Prepare a 1-minute oral introduction of your friend.",
    "Memorize 10 new adjectives from today's lesson.",
    "Think of a short story about you and a friend to share next class."
], 23, color=(45, 52, 54))
add_framed_image(slide, 'books_friends.jpg', 9.2, 1.9, 3.5, 3.0)
add_textbox(slide, 1, 6.1, 11.333, 0.7, "\"A friend in need is a friend indeed.\"", 28, italic=True, color=(251, 194, 235), align=PP_ALIGN.CENTER)
add_textbox(slide, 1, 6.8, 11.333, 0.5, "See you next time! 👋", 20, color=(255, 255, 255), align=PP_ALIGN.CENTER)
add_textbox(slide, 11, 6.8, 1.5, 0.3, "20 / 20", 14, color=(255, 255, 255), align=PP_ALIGN.RIGHT)

prs.save('/Users/pedro/Documents/GitHub/Xinhan/lessons/friendship_presentation.pptx')
print("PPT created successfully with integrated images!")
