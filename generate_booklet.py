#!/usr/bin/env python3
"""Generate Booklet bilingual document from old bilingual docx."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_run_format(run, bold=False, size_pt=None, italic=False):
    run.bold = bold
    run.italic = italic
    if size_pt:
        run.font.size = Pt(size_pt)


def add_para(doc, text, align=None, bold=False, size_pt=None, space_before=None, space_after=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_format(run, bold=bold, size_pt=size_pt)
    return p


def generate_booklet():
    src = Document('uncoveredDocs/hainanDuties2026_05/BookletSpecificatoinsHainan_Bilingual.docx')
    doc = Document()

    # Extract all tables
    tables = src.tables
    print(f'Found {len(tables)} tables')

    # Table 1: Title
    t1 = tables[0]
    left_paras = [p.text.strip() for p in t1.rows[0].cells[0].paragraphs if p.text.strip()]
    right_paras = [p.text.strip() for p in t1.rows[0].cells[1].paragraphs if p.text.strip()]

    cn_title = right_paras[0] if right_paras else ''
    en_title = left_paras[0] if left_paras else ''
    # Clean up English title (keep original text)
    en_title = en_title.replace('  ', ' ').strip()

    add_para(doc, cn_title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)
    add_para(doc, en_title, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size_pt=18)
    doc.add_paragraph()

    # Process remaining tables
    for ti in range(1, len(tables)):
        tbl = tables[ti]
        if len(tbl.rows) == 0:
            continue
        row = tbl.rows[0]
        if len(row.cells) < 2:
            continue

        left_paras = [p.text.strip() for p in row.cells[0].paragraphs if p.text.strip()]
        right_paras = [p.text.strip() for p in row.cells[1].paragraphs if p.text.strip()]

        if not left_paras and not right_paras:
            continue

        # Determine header
        cn_header = right_paras[0] if right_paras else ''
        en_header = left_paras[0] if left_paras else ''

        # Use CN header as the section header (it usually includes English in parentheses)
        header_text = cn_header
        add_para(doc, header_text, bold=True, size_pt=12, space_before=12, space_after=6)

        # Output remaining CN paragraphs
        for ptext in right_paras[1:]:
            if ptext:
                add_para(doc, ptext, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        # Output EN paragraphs
        for ptext in left_paras[1:]:
            if ptext:
                add_para(doc, ptext, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        doc.add_paragraph()

    output_path = 'uncoveredDocs/BookletSpecificatoinsHainan_Bilingual_V20260609.docx'
    doc.save(output_path)
    print(f'Saved: {output_path}')


if __name__ == '__main__':
    generate_booklet()
